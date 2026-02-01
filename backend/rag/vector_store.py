"""
Vector store module for RAG (Retrieval-Augmented Generation).
Supports: PDF, Word (.docx), TXT, PPTX, Images (with OCR)
Uses fallback loaders (pypdf, python-docx, python-pptx) when LangChain loaders fail.
"""
from langchain_openai import OpenAIEmbeddings
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from backend import config
import os

# Optional loaders - use fallbacks if not available
def _load_pdf(filepath):
    try:
        from langchain_community.document_loaders import PyPDFLoader
        loader = PyPDFLoader(filepath)
        return loader.load()
    except Exception:
        try:
            from pypdf import PdfReader
            reader = PdfReader(filepath)
            docs = [Document(page_content=p.extract_text() or "", metadata={"source": filepath, "page": i + 1}) for i, p in enumerate(reader.pages)]
            return docs
        except ImportError:
            raise RuntimeError("`pypdf` package not found, please install it with `pip install pypdf`")

def _load_docx(filepath):
    try:
        from langchain_community.document_loaders import UnstructuredWordDocumentLoader
        loader = UnstructuredWordDocumentLoader(filepath)
        return loader.load()
    except Exception:
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(filepath)
            text = "\n".join(p.text for p in doc.paragraphs)
            return [Document(page_content=text, metadata={"source": filepath})]
        except ImportError:
            raise RuntimeError("unstructured or python-docx not found. Install: `pip install unstructured` or `pip install python-docx`")

def _load_pptx(filepath):
    try:
        from langchain_community.document_loaders import UnstructuredPowerPointLoader
        loader = UnstructuredPowerPointLoader(filepath)
        return loader.load()
    except Exception:
        try:
            from pptx import Presentation
            prs = Presentation(filepath)
            parts = []
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text.append(shape.text)
                parts.append("\n".join(slide_text))
            text = "\n\n".join(parts)
            return [Document(page_content=text, metadata={"source": filepath})]
        except ImportError:
            raise RuntimeError("unstructured or python-pptx not found. Install: `pip install unstructured` or `pip install python-pptx`")

try:
    from langchain_community.document_loaders import TextLoader
except ImportError:
    TextLoader = None

# Try to import FAISS first, fallback to Chroma, then mock
FAISS_AVAILABLE = False
CHROMA_AVAILABLE = False
try:
    from langchain_community.vectorstores import FAISS
    FAISS_AVAILABLE = True
except (ImportError, Exception) as e:
    print(f"[VectorStore] FAISS not available: {e}")
    try:
        from langchain_community.vectorstores import Chroma
        CHROMA_AVAILABLE = True
        print("[VectorStore] Using Chroma as fallback")
    except (ImportError, Exception) as e2:
        print(f"[VectorStore] Chroma not available: {e2}")
        print("[VectorStore] Will use mock vector store")

# Import mock embeddings for test mode
if config.TEST_MODE:
    from backend.rag.mock_embeddings import MockEmbeddings
try:
    from PIL import Image
    import pytesseract
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False
    print("Warning: OCR libraries not installed. Image processing disabled.")

# Chunking configuration with overlap (from config)
CHUNK_SIZE = config.CHUNK_SIZE
CHUNK_OVERLAP = config.CHUNK_OVERLAP

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    length_function=len,
    separators=["\n\n", "\n", " ", ""]
)

# Initialize embeddings - use mock in test mode, OpenAI otherwise
if config.TEST_MODE:
    print("[VectorStore] Using mock embeddings (test mode)")
    embedding_function = MockEmbeddings()
else:
    embedding_function = OpenAIEmbeddings(openai_api_key=config.OPENAI_API_KEY)

# Create a mock vector store class for when Chroma is not available
class MockVectorStore:
    """Mock vector store for when ChromaDB is not available."""
    def __init__(self, persist_directory=None, embedding_function=None):
        self.persist_directory = persist_directory
        self.embedding_function = embedding_function
        self._docs = []
    
    def similarity_search(self, query, k=3):
        """Return empty results (mock)."""
        return []
    
    def add_documents(self, docs):
        """Store documents (mock)."""
        self._docs.extend(docs)
    
    def persist(self):
        """Persist (mock - no-op)."""
        pass

# Initialize vector store - prefer FAISS, fallback to Chroma, then mock
_db = None

def get_db():
    """Get or create the vector store instance. Uses FAISS if available."""
    global _db
    if _db is None:
        if FAISS_AVAILABLE:
            try:
                import os
                index_path = config.FAISS_INDEX_DIR
                if os.path.exists(index_path) and os.path.exists(os.path.join(index_path, "index.faiss")):
                    _db = FAISS.load_local(index_path, embedding_function, allow_dangerous_deserialization=True)
                    print("[VectorStore] Loaded FAISS index from", index_path)
                else:
                    # Create minimal empty index until first ingest
                    _db = FAISS.from_texts([" "], embedding_function)
                    print("[VectorStore] Created new FAISS index (run ingest to populate)")
            except Exception as e:
                print(f"[VectorStore] Could not initialize FAISS: {e}")
                _db = _create_fallback_db()
        else:
            _db = _create_fallback_db()
    return _db

def _create_fallback_db():
    """Create Chroma or Mock fallback."""
    if CHROMA_AVAILABLE:
        try:
            return Chroma(
                persist_directory=config.CHROMA_PERSIST_DIR,
                embedding_function=embedding_function
            )
        except Exception as e:
            print(f"[VectorStore] Chroma init failed: {e}")
    return MockVectorStore(embedding_function=embedding_function)

# For backward compatibility - create a simple wrapper
class DBWrapper:
    def similarity_search(self, *args, **kwargs):
        return get_db().similarity_search(*args, **kwargs)
    def add_documents(self, *args, **kwargs):
        return get_db().add_documents(*args, **kwargs)
    def persist(self, *args, **kwargs):
        return get_db().persist(*args, **kwargs)

db = DBWrapper()

def load_image_with_ocr(image_path):
    """Extract text from images using OCR."""
    if not OCR_AVAILABLE:
        return []
    
    try:
        image = Image.open(image_path)
        text = pytesseract.image_to_string(image)
        
        # Create a document-like structure
        from langchain_core.documents import Document
        return [Document(page_content=text, metadata={"source": image_path, "type": "image"})]
    except Exception as e:
        print(f"Error processing image {image_path}: {e}")
        return []


def ingest_docs(docs_dir: str = None):
    """Ingest all documents from the docs directory with chunking.
    Supports: PDF, Word (.docx), TXT, PPTX, Images (with OCR)
    Returns: dict with chunks_created, files_processed, errors
    """
    docs_dir = docs_dir or config.DOCS_DIR
    all_chunks = []
    files_processed = 0
    errors = []
    
    if not os.path.exists(docs_dir):
        print(f"Directory {docs_dir} does not exist")
        return {"chunks_created": 0, "files_processed": 0, "errors": [f"Directory {docs_dir} does not exist"]}
    
    # Load all supported file types
    for filename in sorted(os.listdir(docs_dir)):
        filepath = os.path.join(docs_dir, filename)
        
        try:
            if filename.endswith(".pdf"):
                docs = _load_pdf(filepath)
                print(f"Loaded PDF {filename}: {len(docs)} pages")
                
            elif filename.endswith((".docx", ".doc")):
                docs = _load_docx(filepath)
                print(f"Loaded Word {filename}: {len(docs)} pages")
                
            elif filename.endswith(".txt"):
                if TextLoader:
                    docs = TextLoader(filepath, encoding='utf-8').load()
                else:
                    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                        docs = [Document(page_content=f.read(), metadata={"source": filepath})]
                print(f"Loaded TXT {filename}: {len(docs)} pages")
                
            elif filename.endswith((".pptx", ".ppt")):
                docs = _load_pptx(filepath)
                print(f"Loaded PowerPoint {filename}: {len(docs)} slides")
                
            elif filename.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp")):
                if OCR_AVAILABLE:
                    docs = load_image_with_ocr(filepath)
                    print(f"Loaded Image {filename} (OCR): {len(docs)} pages")
                else:
                    print(f"Skipping {filename}: OCR not available")
                    continue
            else:
                continue
            
            # Chunk documents with overlap
            chunks = text_splitter.split_documents(docs)
            all_chunks.extend(chunks)
            files_processed += 1
            print(f"  -> {len(chunks)} chunks created")
            
        except Exception as e:
            err_msg = f"Error loading {filename}: {e}"
            print(err_msg)
            errors.append(err_msg)
            continue
    
    if all_chunks:
        global _db
        if FAISS_AVAILABLE:
            os.makedirs(config.FAISS_INDEX_DIR, exist_ok=True)
            _db = FAISS.from_documents(all_chunks, embedding_function)
            _db.save_local(config.FAISS_INDEX_DIR)
            print(f"\nTotal chunks ingested: {len(all_chunks)} (chunk_size={CHUNK_SIZE}, overlap={CHUNK_OVERLAP})")
            print(f"[VectorStore] FAISS index saved to {config.FAISS_INDEX_DIR}")
        else:
            db.add_documents(all_chunks)
            db.persist()
            print(f"\nTotal chunks ingested: {len(all_chunks)} (chunk_size={CHUNK_SIZE}, overlap={CHUNK_OVERLAP})")
    else:
        print("No supported files found to ingest")
    
    return {"chunks_created": len(all_chunks), "files_processed": files_processed, "errors": errors}

